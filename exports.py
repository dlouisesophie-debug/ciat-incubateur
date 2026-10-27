"""Exports Excel (openpyxl) et PDF (reportlab)."""
import datetime as dt
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import app as A

BLUE = colors.HexColor("#3e2214")


def _esc(s):
    return (str(s) if s is not None else "—").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _d(s):
    return A.fr_date(s)


def build_xlsx(db):
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="3E2214")

    def sheet(ws, headers, rows):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
        ws.freeze_panes = "A2"

    projects = db.execute("""SELECT p.id, p.title, o.name owner, p.domain, p.stage, p.cohort, p.status, m.name mentor,
        p.submitted_at, (SELECT ROUND(AVG((e.innovation*25.0+e.faisabilite*20+e.impact*25+e.equipe*15+e.marche*15)/5),1)
                         FROM evaluations e WHERE e.project_id=p.id) score
        FROM projects p JOIN users o ON o.id=p.owner_id LEFT JOIN users m ON m.id=p.mentor_id ORDER BY p.id""").fetchall()
    ws = wb.active
    ws.title = "Projets"
    sheet(ws, ["ID", "Titre", "Porteur", "Domaine", "Stade", "Cohorte", "Statut", "Mentor", "Soumis le", "Note moyenne /100"],
          [(r["id"], r["title"], r["owner"], r["domain"], r["stage"], r["cohort"], A.STATUTS[r["status"]],
            r["mentor"], r["submitted_at"], r["score"]) for r in projects])

    ms = db.execute("""SELECT p.title, m.title mt, m.due_date, m.status, m.completed_at,
        CASE WHEN m.status!='fait' AND m.due_date < date('now') THEN 'Oui' ELSE 'Non' END late
        FROM milestones m JOIN projects p ON p.id=m.project_id ORDER BY p.id, m.due_date""").fetchall()
    sheet(wb.create_sheet("Jalons"), ["Projet", "Jalon", "Échéance", "Statut", "Réalisé le", "En retard"],
          [(r["title"], r["mt"], r["due_date"], A.MILESTONE_STATUS[r["status"]], r["completed_at"], r["late"]) for r in ms])

    evs = db.execute("""SELECT p.title, u.name, e.innovation, e.faisabilite, e.impact, e.equipe, e.marche,
        ROUND((e.innovation*25.0+e.faisabilite*20+e.impact*25+e.equipe*15+e.marche*15)/5,1) note, e.comment
        FROM evaluations e JOIN projects p ON p.id=e.project_id JOIN users u ON u.id=e.evaluator_id
        ORDER BY p.id""").fetchall()
    sheet(wb.create_sheet("Évaluations"),
          ["Projet", "Évaluateur", "Innovation /5", "Faisabilité /5", "Impact /5", "Équipe /5", "Marché /5", "Note /100", "Commentaire"],
          [tuple(r) for r in evs])

    fu = db.execute("""SELECT p.title, f.kind, f.due_date, f.status, f.done_date, f.jobs, f.revenue, f.funding, f.notes
        FROM followups f JOIN projects p ON p.id=f.project_id ORDER BY p.id, f.due_date""").fetchall()
    sheet(wb.create_sheet("Suivis 3-6 mois"),
          ["Projet", "Type", "Échéance", "Statut", "Réalisé le", "Emplois", "Chiffre d'affaires (FCFA)", "Financement (FCFA)", "Notes"],
          [(r["title"], r["kind"].replace("_", " "), r["due_date"], "Fait" if r["status"] == "fait" else "À faire",
            r["done_date"], r["jobs"], r["revenue"], r["funding"], r["notes"]) for r in fu])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _styles():
    ss = getSampleStyleSheet()
    return dict(
        title=ParagraphStyle("t", parent=ss["Title"], textColor=BLUE, fontSize=18, spaceAfter=4),
        h2=ParagraphStyle("h2", parent=ss["Heading2"], textColor=BLUE, fontSize=12.5, spaceBefore=10, spaceAfter=4),
        body=ParagraphStyle("b", parent=ss["BodyText"], fontSize=9.5, leading=13),
        small=ParagraphStyle("s", parent=ss["BodyText"], fontSize=8, textColor=colors.grey))


def _table(data, widths):
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BLUE), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f1ec")])]))
    return t


def _logo():
    path = os.path.join(A.BASE_DIR, "static", "logo-isep.jpg")
    return [Image(path, width=3.4 * cm, height=3.4 * cm * 346 / 564, hAlign="LEFT"), Spacer(1, 6)] if os.path.exists(path) else []


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.grey)
    canvas.drawString(2 * cm, 1.2 * cm, "Incubateur du CIAT - ISEP Amadou Trawaré de Diamniadio - document généré par la plateforme")
    canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
    canvas.restoreState()


def build_project_pdf(db, p, evs, scores, avg, comp, author, show_evals=True):
    s = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=2 * cm,
                            title=f"Compte rendu - {p['title']}", author=author)
    story = _logo() + [Paragraph("Compte rendu de projet incubé", s["title"]),
             Paragraph(f"{_esc(p['title'])} - édité le {dt.date.today().strftime('%d/%m/%Y')} par {_esc(author)}", s["small"]),
             Spacer(1, 8), Paragraph("1. Identification", s["h2"])]
    rows = [["Porteur", _esc(p["owner_name"])], ["Domaine", _esc(p["domain"])], ["Stade", _esc(p["stage"])],
            ["Cohorte", _esc(p["cohort"])], ["Statut", A.STATUTS[p["status"]]], ["Mentor", _esc(p["mentor_name"])],
            ["Dossier soumis le", _d(p["submitted_at"])]]
    t = Table([[Paragraph(f"<b>{a}</b>", s["body"]), Paragraph(b, s["body"])] for a, b in rows], colWidths=[4 * cm, 13 * cm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [t, Paragraph("2. Présentation du projet", s["h2"])]
    for lab, key in (("Résumé", "summary"), ("Problème", "problem"), ("Solution", "solution"), ("Public cible", "target")):
        story.append(Paragraph(f"<b>{lab} :</b> {_esc(p[key])}", s["body"]))
    story.append(Paragraph(f"Complétude du dossier : {comp['pct']} %", s["body"]))

    if show_evals:
        story.append(Paragraph("3. Évaluation du comité", s["h2"]))
        if evs:
            data = [["Évaluateur", *[l.split()[0] for _k, l, _w in A.CRITERES], "Note /100"]]
            for e, sc in zip(evs, scores):
                data.append([_esc(e["evaluator"]), *[e[k] for k, _l, _w in A.CRITERES], sc])
            story += [_table(data, [4.2 * cm] + [2.2 * cm] * 5 + [2 * cm]),
                      Paragraph(f"Note moyenne : <b>{avg} / 100</b> ({len(evs)} évaluations)", s["body"])]
        else:
            story.append(Paragraph("Aucune évaluation enregistrée.", s["body"]))
    d = db.execute("""SELECT d.*, u.name by FROM decisions d JOIN users u ON u.id=d.decided_by
                      WHERE project_id=? ORDER BY d.id DESC LIMIT 1""", (p["id"],)).fetchone()
    story.append(Paragraph("4. Décision" if show_evals else "3. Décision", s["h2"]))
    if d:
        lab = {"accepte": "Accepté en incubation", "refuse": "Refusé", "liste_attente": "Liste d'attente"}[d["decision"]]
        story.append(Paragraph(f"<b>{lab}</b> le {_d(d['decided_at'])} (par {_esc(d['by'])}).<br/>Motivation : {_esc(d['motivation'])}", s["body"]))
    else:
        story.append(Paragraph("Aucune décision enregistrée.", s["body"]))

    objs = db.execute("SELECT * FROM objectives WHERE project_id=? ORDER BY due_date", (p["id"],)).fetchall()
    story.append(Paragraph("Objectifs", s["h2"]))
    if objs:
        story.append(_table([["Objectif", "Indicateur", "Échéance"]] +
                            [[Paragraph(_esc(o["label"]), s["body"]), Paragraph(_esc(o["indicator"]), s["body"]), _d(o["due_date"])] for o in objs],
                            [8 * cm, 6 * cm, 3 * cm]))
    else:
        story.append(Paragraph("Aucun objectif défini.", s["body"]))

    ms = db.execute("SELECT * FROM milestones WHERE project_id=? ORDER BY due_date", (p["id"],)).fetchall()
    story.append(Paragraph("Jalons", s["h2"]))
    if ms:
        rows = [["Jalon", "Échéance", "Statut", "Preuve"]]
        for m in ms:
            late = m["status"] != "fait" and m["due_date"] < dt.date.today().isoformat()
            rows.append([Paragraph(_esc(m["title"]), s["body"]), _d(m["due_date"]),
                         A.MILESTONE_STATUS[m["status"]] + (" - EN RETARD" if late else ""), "Oui" if m["proof_stored"] else "Non"])
        story.append(_table(rows, [7.5 * cm, 2.5 * cm, 4.5 * cm, 2.5 * cm]))
        done = sum(m["status"] == "fait" for m in ms)
        story.append(Paragraph(f"Avancement : {done}/{len(ms)} jalons réalisés ({round(done / len(ms) * 100)} %).", s["body"]))
    else:
        story.append(Paragraph("Aucun jalon défini.", s["body"]))

    fu = db.execute("SELECT * FROM followups WHERE project_id=? ORDER BY due_date", (p["id"],)).fetchall()
    if fu:
        story.append(Paragraph("Suivi à 3 et 6 mois", s["h2"]))
        story.append(_table([["Suivi", "Échéance", "Statut", "Emplois", "CA (FCFA)", "Financement (FCFA)"]] +
                            [[f["kind"].replace("_", " "), _d(f["due_date"]), "Fait" if f["status"] == "fait" else "À faire",
                              f["jobs"] if f["jobs"] is not None else "—", f["revenue"] if f["revenue"] is not None else "—",
                              f["funding"] if f["funding"] is not None else "—"] for f in fu],
                            [2.4 * cm, 2.6 * cm, 2.4 * cm, 2.2 * cm, 3.2 * cm, 4.2 * cm]))
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def build_list_pdf(db, d):
    s = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=2 * cm,
                            title="Synthèse de l'incubateur du CIAT")
    story = _logo() + [Paragraph("Synthèse de l'incubateur du CIAT", s["title"]),
             Paragraph(f"Édité le {dt.date.today().strftime('%d/%m/%Y')}", s["small"]), Paragraph("Indicateurs clés", s["h2"])]
    kpis = [["Indicateur", "Valeur"],
            ["Projets déposés (total)", d["total"]], ["Dossiers soumis au comité", d["submitted"]],
            ["Projets acceptés en incubation", d["accepted"]],
            ["Taux d'acceptation (accepté / décidé)", f"{d['acceptance_rate']} %" if d["acceptance_rate"] is not None else "—"],
            ["Note moyenne des évaluations", f"{d['avg']} / 100" if d["avg"] is not None else "—"],
            ["Jalons réalisés", f"{d['ms']['done'] or 0} / {d['ms']['total'] or 0}"], ["Jalons en retard", d["ms"]["late"] or 0]]
    story += [_table(kpis, [11 * cm, 6 * cm]), Paragraph("Répartition par statut", s["h2"])]
    story.append(_table([["Statut", "Nombre"]] + [[A.STATUTS[k], v] for k, v in d["by_status"].items()], [11 * cm, 6 * cm]))
    story.append(Paragraph("Répartition par domaine", s["h2"]))
    story.append(_table([["Domaine", "Projets"]] + [[r["d"], r["c"]] for r in d["by_domain"]], [11 * cm, 6 * cm]))
    story.append(Paragraph("Comparaison des cohortes", s["h2"]))
    story.append(_table([["Cohorte", "Projets", "Acceptés"]] + [[r["k"], r["total"], r["accepted"] or 0] for r in d["by_cohort"]],
                        [8 * cm, 4.5 * cm, 4.5 * cm]))
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()
